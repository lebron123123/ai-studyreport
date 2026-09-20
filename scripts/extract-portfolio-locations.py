"""Extract location evidence only; preserve source workbook and financial ledger."""
import sys, json, hashlib
from pathlib import Path
from openpyxl import load_workbook

source, target = map(Path, sys.argv[1:3])
sheet = load_workbook(source, data_only=True)['拓展项目清单']
assert sheet.cell(2, 4).value == '具体地址'
items = []
for row in range(3, sheet.max_row + 1):
    name = sheet.cell(row, 3).value
    if not name or str(name).strip() in ('合计', '总计'):
        continue
    name = str(name).strip()
    address = str(sheet.cell(row, 4).value or '').strip()
    coordinate = str(sheet.cell(row, 7).value or '').strip()
    point = None
    try:
        lat, lng = map(float, coordinate.replace('，', ',').split(','))
        if 20 < lat < 25 and 110 < lng < 117:
            point = [lng, lat]
    except ValueError:
        pass
    items.append(dict(name=name, sourceKey=hashlib.sha256(('anju-ledger\0'+name).encode()).hexdigest()[:32],
                      address='' if address == '/' else address, suppliedPoint=point,
                      suppliedCoordinate=coordinate, suppliedCrs='unknown', coordinateSource='用户说明：豆包搜索', row=row))
assert len({p['sourceKey'] for p in items}) == len(items)
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(dict(format='anju-location-evidence-v1', sourceFile=source.name,
    sourceHash=hashlib.sha256(source.read_bytes()).hexdigest(), projects=items), ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(dict(total=len(items), coordinates=sum(p['suppliedPoint'] is not None for p in items),
                     missing=[p['name'] for p in items if p['suppliedPoint'] is None]), ensure_ascii=False))
