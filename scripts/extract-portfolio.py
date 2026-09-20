"""Read-only extraction of the supplied ledger; never changes the workbook."""
import sys, json, hashlib, datetime
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

source, target = map(Path, sys.argv[1:3])
book = load_workbook(source, data_only=True)
formulas = load_workbook(source, data_only=False)
projects = {}
for sheet in book:
    owner = ''
    for row in range(3, sheet.max_row + 1):
        name = sheet.cell(row, 3).value
        serial = sheet.cell(row, 2).value
        if sheet.cell(row, 1).value:
            owner = str(sheet.cell(row, 1).value)
        if not name or str(name).strip() in ('合计', '总计'):
            continue
        name = str(name).strip()
        if not isinstance(serial, (int, float)):
            print('CHECK unnumbered row', sheet.title, row, name)
        key = name
        item = projects.setdefault(key, dict(name=name, category='completed', potentialTier='', address='', point=None,
            locationStatus='pending', owner=owner, type='', records=[]))
        fields = []
        for col in range(1, sheet.max_column + 1):
            cell = sheet.cell(row, col)
            value = cell.value
            label = sheet.cell(2, col).value
            if label is None and value is None:
                continue
            if col == 1 and value is None:
                value = owner
            if col in (9, 10) and isinstance(value, (int, float)) and 20000 < value < 80000:
                value = from_excel(value, book.epoch)
            if isinstance(value, (datetime.datetime, datetime.date)):
                value = value.isoformat()[:10]
            display = '未填' if value is None else str(value)
            if isinstance(value, (int, float)) and '%' in cell.number_format:
                display = f'{value * 100:.2f}%'
            formula = formulas[sheet.title].cell(row, col).value
            fields.append(dict(label=str(label or '辅助列').replace('\n', ' '), value=value, display=display,
                cell=cell.coordinate, numberFormat=cell.number_format, formula=formula if isinstance(formula, str) and formula.startswith('=') else ''))
        item['records'].append(dict(sheet=sheet.title, row=row, fields=fields))
        item['type'] = str(sheet.cell(row, 4).value or item['type'])
for item in projects.values():
    item['sourceKey'] = hashlib.sha256(('anju-ledger\0' + item['name']).encode()).hexdigest()[:32]
    item['sourceFile'] = source.name
    item['sourceHash'] = hashlib.sha256(source.read_bytes()).hexdigest()
    item['preferredRecord'] = len(item['records']) - 1
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(dict(format='anju-portfolio-v1', projects=list(projects.values())), ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(dict(projects=len(projects), records=sum(len(p['records']) for p in projects.values()), output=str(target)), ensure_ascii=False))
